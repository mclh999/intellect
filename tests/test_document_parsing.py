"""T2.1/T2.2：真实格式解析、固定统计、证据切分及批次恢复。"""

from __future__ import annotations

import io
import json
from contextlib import closing
from email.message import Message
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, DecodedStreamObject, NameObject, NumberObject

from intellect.core.config import ChunkConfig, Settings, TokenizationConfig
from intellect.db import connect_database, initialize_database
from intellect.doc import counting, hashing, parsers, processing, web
from intellect.doc.counting import ChunkTokenCounter, TokenCountError, get_counter
from intellect.doc.hashing import DocumentHashError, chunk_hash, document_hash, file_hash
from intellect.doc.parsers import DocumentParseError, inspect_bytes, parse_bytes, parser_version
from intellect.doc.processing import DocumentProcessingError, DocumentProcessingService
from intellect.doc.splitting import split_document
from intellect.search.tokenization import TextTokenizer


def pdf_bytes(texts=("First page", "Second page"), *, encrypted=False, image=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 30 250 Td ({text}) Tj ET".encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
        if image:
            stream = DecodedStreamObject()
            stream.set_data(b'\xff\x00\x00')
            stream.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                           NameObject('/Width'): NumberObject(1), NameObject('/Height'): NumberObject(1),
                           NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'): DictionaryObject({NameObject('/Im1'): writer._add_object(stream)})})
    if encrypted:
        writer.encrypt('password')
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    document = Document()
    paragraph = document.add_paragraph("第一段\t不可以")
    paragraph.add_run().add_break()
    paragraph.add_run("10 mg")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "左单元格"
    table.cell(0, 1).text = "右单元格"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表格"
    document.add_paragraph("结尾段落")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def connection(tmp_path):
    path = tmp_path / 'parsing.db'
    initialize_database(path)
    with closing(connect_database(path)) as database:
        yield database


def registered(connection, tmp_path, data=b'Alpha Beta Gamma Delta Epsilon', kind='TXT'):
    path = tmp_path / ('archive.' + {'TXT':'txt','PDF':'pdf','DOCX':'docx','WEB':'html'}[kind])
    path.write_bytes(data)
    content = inspect_bytes(data, kind)
    with connection:
        document_id = connection.execute(
            "INSERT INTO source_document (title, storage_url, source_type, file_hash, content_hash, file_size) "
            "VALUES ('原件', ?, ?, ?, ?, ?)",
            (str(path), kind, content.file_hash, content.content_hash, len(data)),
        ).lastrowid
    return document_id, path


class TestDocumentHashes:
    @pytest.mark.parametrize('function,input_value',[(file_hash,b'abc'),(document_hash,'abc')])
    def test_protocol_manifest_mismatch_fails_closed(self,tmp_path,monkeypatch,function,input_value):
        (tmp_path/'hash_protocols.json').write_text('{}',encoding='utf-8')
        hashing._document_prefix.cache_clear()
        monkeypatch.setattr(hashing,'files',lambda package:tmp_path)
        try:
            with pytest.raises(DocumentHashError):
                function(input_value)
        finally:
            hashing._document_prefix.cache_clear()

    def test_fixed_digest_and_only_document_newline_normalization(self):
        expected = 'ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad'
        assert file_hash(b'abc') == 'file-v1:sha256:' + expected
        assert document_hash('abc') == 'document-v1:sha256:' + expected
        assert document_hash('a\r\nb\rc') == document_hash('a\nb\nc')
        assert chunk_hash('a\r\nb') != chunk_hash('a\nb')
        assert file_hash(b'a\r\nb') != file_hash(b'a\nb')

    @pytest.mark.parametrize('left,right', [('可以','不可以'),('10 mg','100 mg'),('a b','ab'),('é','e\u0301'),('ABC','abc'),(' A','A'),('A.','A'),('\ufeffA','A')])
    def test_meaning_and_evidence_differences_preserved(self, left, right):
        assert document_hash(left) != document_hash(right)

    def test_invalid_hash_inputs_fail(self):
        for call, value in ((file_hash,'abc'),(document_hash,b'abc'),(document_hash,'\ud800')):
            with pytest.raises(DocumentHashError):
                call(value)


class TestFormatParsers:
    def test_txt_preserves_bom_tabs_whitespace_and_newlines(self):
        text = '\ufeff  第一行\r\n\t第二行\r第三行\n'
        result = parse_bytes(text.encode('utf-8'),'TXT')
        assert result.content == text
        assert result.blocks[0].start == 0

    def test_explicit_encoding_and_strict_decode(self):
        data = '中文'.encode('gb18030')
        assert parse_bytes(data,'TXT',options={'encoding':'gb18030'}).content == '中文'
        with pytest.raises(DocumentParseError):
            parse_bytes(data,'TXT')
        with pytest.raises(DocumentParseError):
            parse_bytes(b'abc','TXT',options={'encoding':'unknown-encoding'})

    def test_pdf_real_text_and_page_metadata(self):
        result = parse_bytes(pdf_bytes(),'PDF')
        assert 'First page' in result.content and 'Second page' in result.content
        assert [block.location['page'] for block in result.blocks] == [1,2]
        assert ''.join(block.text for block in result.blocks) == result.content

    def test_pdf_empty_encrypted_and_image_only_paths(self):
        assert parse_bytes(pdf_bytes(('',)),'PDF').content == ''
        with pytest.raises(DocumentParseError,match='加密'):
            parse_bytes(pdf_bytes(encrypted=True),'PDF')
        with pytest.raises(DocumentParseError,match='OCR'):
            parse_bytes(pdf_bytes(('',),image=True),'PDF')

    def test_docx_paragraphs_tables_and_nested_locations_in_order(self):
        result = parse_bytes(docx_bytes(),'DOCX')
        texts = ['第一段\t不可以\n10 mg','左单元格','右单元格','嵌套表格','结尾段落']
        assert [result.content.index(value) for value in texts] == sorted(result.content.index(value) for value in texts)
        assert any(len(block.location.get('table_path',[])) == 2 for block in result.blocks if '嵌套表格' in block.text)
        for block in result.blocks:
            assert result.content[block.start:block.start+len(block.text)] == block.text

    def test_html_inline_context_entities_and_nontext_elements(self):
        html = '<style>hidden-style</style><p>不<em>可以</em> &amp; 10 mg<br>条件保留</p><script>hidden-script</script><template>inert</template>'
        result = parse_bytes(html.encode(),'WEB')
        assert result.content == '不可以 & 10 mg\n条件保留\n'
        assert 'hidden' not in result.content and 'inert' not in result.content

    @pytest.mark.parametrize('kind,data', [('PDF',b'broken'),('DOCX',b'broken'),('TXT',b'\xff'),('WEB',b'\xff')])
    def test_corrupt_documents_are_explicit_failures(self,kind,data):
        with pytest.raises(DocumentParseError):
            parse_bytes(data,kind)

    def test_unsupported_format_and_options_are_rejected(self):
        for kind, options in [('DOC',{}),('PDF',{'password':'x'}),('TXT',{}),('TXT',{'encoding':'utf-8','extra':True})]:
            with pytest.raises(DocumentParseError):
                parse_bytes(b'content',kind,options=options)


class TestFixedCountingAndSplitting:
    def test_fts_user_dictionary_cannot_change_fixed_counts(self,tmp_path):
        path = tmp_path/'domain.txt'
        path.write_text('智识回溯协议 1000000 n\n',encoding='utf-8')
        fts = TextTokenizer(TokenizationConfig(user_dictionaries=[str(path)]))
        assert fts.tokenize('智识回溯协议') == ['智识回溯协议']
        assert get_counter().count('智识回溯协议') > 1

    def test_fixed_counter_and_manifest(self):
        counter = get_counter()
        assert counter.count('Alpha Beta,\r\n123') == 4
        assert counter.count(' \n\t') == 0
        assert counter.manifest['HMM'] is False
        changed = counter.manifest
        changed['HMM'] = True
        assert counter.manifest['HMM'] is False

    def test_counter_rejects_different_runtime_or_resource(self, monkeypatch, tmp_path):
        monkeypatch.setattr(counting.jieba,'__version__','different')
        with pytest.raises(TokenCountError):
            ChunkTokenCounter()
        monkeypatch.undo()
        (tmp_path/'assets').mkdir()
        (tmp_path/'assets'/'token_counter.json').write_text('{}',encoding='utf-8')
        (tmp_path/'assets'/'jieba-0.42.1.dict.txt').write_bytes(b'changed')
        monkeypatch.setattr(counting,'files',lambda package:tmp_path)
        with pytest.raises(TokenCountError):
            ChunkTokenCounter()

    @pytest.mark.parametrize('size,overlap,separators', [(8,2,['']), (8,0,['\n',' ']), (3,2,['\n']), (8,3,['。',' ']), (7,1,[])])
    def test_original_coverage_offsets_max_size_and_progress(self,size,overlap,separators):
        text = 'Alpha\r\n中文。Beta Gamma\tDelta\r尾部 '
        document = parse_bytes(text.encode(),'TXT')
        config = ChunkConfig(chunk_size=size,chunk_overlap=overlap,separators=separators)
        chunks = list(split_document(document,config,get_counter()))
        covered = set()
        for index, chunk in enumerate(chunks):
            a,b = chunk.location_meta['text_start'],chunk.location_meta['text_end']
            assert chunk.chunk_index == index and 0 < len(chunk.content) <= size
            assert chunk.content == text[a:b]
            assert chunk.token_count == get_counter().count(chunk.content)
            covered.update(range(a,b))
        assert covered == set(range(len(text)))
        assert all(chunks[i+1].location_meta['text_start'] > chunks[i].location_meta['text_start'] for i in range(len(chunks)-1))

    def test_separator_priority_and_structural_page_boundaries(self):
        document = parse_bytes(b'AA BB\nCC DD EE','TXT')
        chunks = list(split_document(document,ChunkConfig(chunk_size=10,chunk_overlap=0,separators=['\n',' ']),get_counter()))
        assert chunks[0].content == 'AA BB\n'
        pages = parse_bytes(pdf_bytes(),'PDF')
        chunks = list(split_document(pages,ChunkConfig(chunk_size=100,chunk_overlap=10),get_counter()))
        assert [chunk.location_meta['page'] for chunk in chunks] == [1,2]


class TestProcessingPipeline:
    def test_unicode_space_and_hash_in_archive_file_uri(self,connection,tmp_path):
        doc,path = registered(connection,tmp_path)
        special = tmp_path/'中文 # 归档.txt'
        special.write_bytes(path.read_bytes())
        with connection:
            connection.execute('UPDATE source_document SET storage_url=? WHERE id=?',(special.as_uri(),doc))
        service = DocumentProcessingService(connection)
        batch = service.create(doc)
        assert '%23' in json.loads(batch.parse_config)['parser_options']['archive']
        assert service.run(batch.task_id).chunk_count > 0

    @pytest.mark.parametrize('kind', ['TXT','PDF','DOCX','WEB'])
    def test_four_formats_publish_valid_hash_counts_and_locations(self,connection,tmp_path,kind):
        data = {'TXT':b'Alpha\r\nBeta','PDF':pdf_bytes(),'DOCX':docx_bytes(),'WEB':b'<p>Alpha</p><p>Beta</p>'}[kind]
        doc,path = registered(connection,tmp_path,data,kind)
        service = DocumentProcessingService(connection)
        batch = service.create(doc,chunking=ChunkConfig(chunk_size=12,chunk_overlap=2))
        snapshot = json.loads(batch.parse_config)
        assert snapshot['parser_options']['file_hash'] == file_hash(data)
        assert snapshot['parser_version'] == parser_version(kind)
        assert snapshot['parser_options']['token_counter'] == get_counter().manifest
        assert connection.execute('SELECT chunk_batch,parse_config FROM source_document').fetchone() == (None,None)
        assert connection.execute('SELECT count(*) FROM document_chunk').fetchone()[0] == 0
        published = service.run(batch.task_id)
        rows = connection.execute('SELECT content,chunk_hash,token_count,location_meta FROM document_chunk ORDER BY chunk_index').fetchall()
        assert published.chunk_count == len(rows) > 0
        for text,digest,count,location in rows:
            assert digest == chunk_hash(text) and count == get_counter().count(text)
            assert json.loads(location)['format'] == kind
        assert path.read_bytes() == data
        assert connection.execute('SELECT chunk_batch,parse_config FROM source_document').fetchone() == (1,batch.parse_config)
        assert connection.execute('SELECT count(*) FROM candidate').fetchone()[0] == 0

    def test_failure_resume_uses_snapshot_and_preserves_chunk_ids(self,connection,tmp_path,monkeypatch):
        doc,_ = registered(connection,tmp_path,'Alpha 中文 Beta Gamma Delta Epsilon'.encode())
        service = DocumentProcessingService(connection)
        batch = service.create(doc,chunking=ChunkConfig(chunk_size=8,chunk_overlap=2,separators=['']))
        original = processing.split_document
        def fail_after_first(*args):
            for index,chunk in enumerate(original(*args)):
                if index == 1:
                    raise RuntimeError('中途失败')
                yield chunk
        monkeypatch.setattr(processing,'split_document',fail_after_first)
        with pytest.raises(RuntimeError):
            service.run(batch.task_id)
        old = connection.execute('SELECT id,content FROM document_chunk').fetchall()
        assert len(old) == 1 and service.batches.get(batch.task_id).chunk_count is None
        assert connection.execute('SELECT chunk_batch FROM source_document').fetchone()[0] is None
        monkeypatch.setattr(processing,'split_document',original)
        monkeypatch.setattr(parsers,'get_settings',lambda:Settings(document={'text_encoding':'ascii'},chunk={'chunk_size':99,'chunk_overlap':0}))
        assert service.run(batch.task_id).chunk_count > 1
        assert connection.execute('SELECT id,content FROM document_chunk ORDER BY chunk_index LIMIT 1').fetchall() == old
        assert service.batches.get(batch.task_id).parse_config == batch.parse_config

    def test_reparse_preserves_old_batch_and_late_retry_cannot_rewind(self,connection,tmp_path,monkeypatch):
        doc,_ = registered(connection,tmp_path)
        service = DocumentProcessingService(connection)
        first = service.run(service.create(doc,chunking=ChunkConfig(chunk_size=9,chunk_overlap=1)).task_id)
        old = connection.execute('SELECT * FROM document_chunk ORDER BY id').fetchall()
        second = service.create(doc,chunking=ChunkConfig(chunk_size=12,chunk_overlap=0))
        original = processing.parse_bytes
        monkeypatch.setattr(processing,'parse_bytes',lambda *a,**k:(_ for _ in ()).throw(DocumentParseError('失败')))
        with pytest.raises(DocumentParseError):
            service.run(second.task_id)
        assert connection.execute('SELECT chunk_batch FROM source_document').fetchone()[0] == 1
        monkeypatch.setattr(processing,'parse_bytes',original)
        service.run(second.task_id)
        assert connection.execute('SELECT * FROM document_chunk WHERE chunk_batch=1 ORDER BY id').fetchall() == old
        service.run(first.task_id)
        assert connection.execute('SELECT chunk_batch FROM source_document').fetchone()[0] == 2

    def test_archive_changes_and_parser_versions_do_not_mix_with_old_batch(self,connection,tmp_path,monkeypatch):
        doc,path = registered(connection,tmp_path)
        service = DocumentProcessingService(connection)
        batch = service.create(doc)
        original = path.read_bytes()
        path.write_bytes(b'changed')
        with pytest.raises(DocumentProcessingError,match='字节已改变'):
            service.run(batch.task_id)
        path.write_bytes(original)
        monkeypatch.setattr(processing,'parser_version',lambda kind:'new-version')
        with pytest.raises(DocumentProcessingError,match='快照'):
            service.run(batch.task_id)
        assert connection.execute('SELECT count(*) FROM document_chunk').fetchone()[0] == 0

    def test_published_task_does_not_reparse(self,connection,tmp_path,monkeypatch):
        doc,_ = registered(connection,tmp_path)
        service = DocumentProcessingService(connection)
        published = service.run(service.create(doc).task_id)
        monkeypatch.setattr(processing,'parse_bytes',lambda *a,**k:(_ for _ in ()).throw(AssertionError('不能重新解析')))
        assert service.run(published.task_id) == published

    @pytest.mark.parametrize('field',['token_counter','splitter_version'])
    def test_snapshot_version_mismatch_prevents_new_writes(self,connection,tmp_path,field):
        doc,_ = registered(connection,tmp_path)
        service = DocumentProcessingService(connection)
        batch = service.create(doc)
        snapshot = json.loads(batch.parse_config)
        snapshot['parser_options'][field] = 'different'
        with connection:
            connection.execute('UPDATE task_process SET parse_config=? WHERE id=?',(json.dumps(snapshot),batch.task_id))
        with pytest.raises(DocumentProcessingError,match='快照'):
            service.run(batch.task_id)
        assert connection.execute('SELECT count(*) FROM document_chunk').fetchone()[0] == 0

    def test_empty_document_publishes_zero(self,connection,tmp_path):
        doc,_ = registered(connection,tmp_path,b'')
        service = DocumentProcessingService(connection)
        result = service.run(service.create(doc).task_id)
        assert result.chunk_count == 0

    def test_outer_transaction_controls_all_stage_writes(self,connection,tmp_path):
        doc,_ = registered(connection,tmp_path)
        connection.execute('BEGIN IMMEDIATE')
        service = DocumentProcessingService(connection)
        service.run(service.create(doc).task_id)
        assert connection.in_transaction
        connection.rollback()
        assert connection.execute('SELECT count(*) FROM task_process').fetchone()[0] == 0
        assert connection.execute('SELECT count(*) FROM document_chunk').fetchone()[0] == 0
        assert connection.execute('SELECT chunk_batch FROM source_document').fetchone()[0] is None

    def test_missing_and_unarchived_sources_do_not_allocate_tasks(self,connection,tmp_path):
        doc,path = registered(connection,tmp_path)
        with connection:
            connection.execute("UPDATE source_document SET storage_url='https://example.test/source' WHERE id=?",(doc,))
        service = DocumentProcessingService(connection)
        with pytest.raises(DocumentProcessingError,match='先归档'):
            service.create(doc)
        with pytest.raises(DocumentProcessingError):
            service.create(999)
        assert connection.execute('SELECT count(*) FROM task_process').fetchone()[0] == 0


class TestWebArchive:
    def test_http_archive_then_processing_works_after_server_stops(self,connection,tmp_path):
        data = '<p>归档后的中文正文</p>'.encode()
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args):
                pass
            def do_GET(self):
                self.send_response(200)
                self.send_header('Content-Type','text/html; charset=utf-8')
                self.end_headers()
                self.wfile.write(data)
        server = ThreadingHTTPServer(('127.0.0.1',0),Handler)
        thread = Thread(target=lambda:server.serve_forever(poll_interval=0.01),daemon=True)
        thread.start()
        try:
            archive = web.archive_web(f'http://127.0.0.1:{server.server_port}/page',tmp_path/'archived.html')
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)
        inspected = inspect_bytes(data,'WEB')
        with connection:
            doc = connection.execute("INSERT INTO source_document(title,storage_url,source_type,file_hash,content_hash) VALUES ('WEB',?,'WEB',?,?)",(str(archive.path),archive.file_hash,inspected.content_hash)).lastrowid
        service = DocumentProcessingService(connection)
        assert service.run(service.create(doc).task_id).chunk_count == 1
        assert archive.path.read_bytes() == data

    def test_response_bytes_archived_once_and_used_for_file_hash(self,tmp_path,monkeypatch):
        data = '<p>中文</p>'.encode('gb18030')
        headers = Message()
        headers['Content-Type'] = 'text/html; charset=gb18030'
        class Response(io.BytesIO):
            def geturl(self):
                return 'https://example.test/final'
        response = Response(data)
        response.headers = headers
        calls = []
        def fetch(request,**kwargs):
            calls.append((request.full_url,kwargs))
            return response
        monkeypatch.setattr(web,'urlopen',fetch)
        monkeypatch.setattr(web,'get_settings',lambda:Settings(document={'web_timeout_seconds':7}))
        archive = web.archive_web('https://example.test/start',tmp_path/'page.html')
        assert archive.path.read_bytes() == data and archive.file_hash == file_hash(data)
        assert archive.encoding == 'gb18030' and archive.source_url.endswith('/final')
        assert parse_bytes(data,'WEB',options={'encoding':archive.encoding}).content == '中文\n'
        with pytest.raises(DocumentParseError,match='禁止覆盖'):
            web.archive_web('https://example.test/start',archive.path)
        assert len(calls) == 1 and calls[0][1]['timeout'] == 7

    def test_invalid_scheme_and_failed_download_leave_no_archive(self,tmp_path,monkeypatch):
        path = tmp_path/'page.html'
        with pytest.raises(DocumentParseError):
            web.archive_web('file:///local',path)
        monkeypatch.setattr(web,'urlopen',lambda *a,**k:(_ for _ in ()).throw(OSError('network failed')))
        with pytest.raises(DocumentParseError):
            web.archive_web('https://example.test/start',path)
        assert not path.exists()
